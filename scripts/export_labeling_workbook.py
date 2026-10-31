"""Create a blind, private Excel workbook for FRIDAY email labeling.

Run from the repository root. Source data is read from data/email_labels.csv;
old human labels are used only to select a mode's rows and are never written to
any workbook sheet, cell, comment, formula, or workbook metadata.
"""
from __future__ import annotations

import argparse
import csv
import json
import random
import re
import shutil
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName

from scripts.label_emails import _fully_labeled, _load_queue_ids
from scripts.sample_test_v2 import has_any_human_label, read_frozen_test_ids

PROJECT_ROOT = Path.cwd()
DEFAULT_INPUT = PROJECT_ROOT / "data" / "email_labels.csv"
DEFAULT_OUTPUT_DIR = PROJECT_ROOT / "data" / "labeling"
DEFAULT_QUEUE = PROJECT_ROOT / "data" / "labeling_queue_minority.csv"
DEFAULT_SPLIT_MANIFEST = PROJECT_ROOT / "data" / "priority_engine_split.json"
DEFAULT_TEST_MANIFEST = PROJECT_ROOT / "data" / "test_v2_manifest.json"
SCHEMA_VERSION = "v2"
GUIDE_VERSION = "v2.0-draft"
DEFAULT_SEED = 20261031
DEFAULT_BATCH_SIZE = 40
DEFAULT_BODY_PREVIEW_CHARS = 1500
MAX_CELL_CHARS = 32767

YES_NO = ("yes", "no")
PRIORITIES = ("URGENT", "NEEDS_REPLY", "INFORMATIONAL")
CATEGORIES = ("JOB", "FINANCE", "PERSONAL", "PROMOTIONS", "UPDATES", "OTHER")
LABEL_COLUMNS = ("is_spam", "needs_response", "priority", "category", "unsure", "note")
VISIBLE_COLUMNS = ("row_number", "email_id", "sender", "subject", "snippet", "body_preview", *LABEL_COLUMNS)
FORBIDDEN_COLUMN_PREFIXES = ("true_", "final_", "predicted_", "rules_", "llm_")
FORBIDDEN_EXACT = {"label_source"}

XML_ILLEGAL = re.compile("[\\x00-\\x08\\x0B\\x0C\\x0E-\\x1F\\uFFFE\\uFFFF]")

GUIDE_ROWS = [
    ("Schema", "FRIDAY email labeling schema v2; guide v2.0-draft (20-email pilot draft)."),
    ("is_spam=yes", "Unwanted, deceptive, or unsolicited content, including phishing, impersonation, scams, malicious links/attachments, and irrelevant unsolicited bulk mail."),
    ("is_spam=no", "Legitimate correspondence, expected service notices, and marketing from an organization or mailing list the recipient signed up for. Automation alone is not spam."),
    ("needs_response=yes", "A real person or organization is genuinely waiting for a reply, confirmation, decision, requested information, or action from you."),
    ("needs_response=no", "No response or action is expected. The fact that you could optionally reply, click, or read does not make it required."),
    ("URGENT", "A genuinely time-sensitive or high-impact matter requiring prompt attention. Look for an imminent event, stated deadline, active incident, or material consequence; ignore dramatic wording alone."),
    ("NEEDS_REPLY", "A reply, confirmation, decision, or requested information is expected, but immediate handling is not established."),
    ("INFORMATIONAL", "No response is expected and no separate urgent action is established."),
    ("JOB", "Recruitment, job openings, internships, applications, interviews, assessments, offers, and employment opportunities."),
    ("FINANCE", "Banking, payments, bills, invoices, taxes, transactions, and financial services."),
    ("PERSONAL", "Direct communication involving friends, family, or personal acquaintances, including personal arrangements."),
    ("PROMOTIONS", "Legitimate marketing, discounts, sales, product offers, and promotional newsletters."),
    ("UPDATES", "Service, account, delivery, subscription-status, product/system, and transaction-status notifications whose primary purpose is to inform."),
    ("OTHER", "A relevant email that does not fit the other five categories. Do not use OTHER as a substitute for uncertainty."),
    ("Decision order", "Read context; decide spam; decide whether a response/action is expected; decide urgency; assign the primary category; mark uncertainty honestly."),
    ("Consistency", "needs_response=yes cannot pair with INFORMATIONAL. needs_response=no cannot pair with NEEDS_REPLY. URGENT may pair with needs_response=no when immediate action is needed but no reply is expected. Flag conflicts; never auto-correct."),
    ("Spam vs PROMOTIONS", "Subscribed, legitimate marketing = is_spam=no and PROMOTIONS. Deceptive, unwanted, or unsolicited bulk content = is_spam=yes."),
    ("Spam/PROMOTIONS examples", "Synthetic: a subscribed bookstore discount newsletter = is_spam=no, PROMOTIONS. A fake prize email asking for a bank password = is_spam=yes."),
    ("UPDATES vs PROMOTIONS", "Existing-service/account/status information = UPDATES. Persuasion to buy, subscribe, upgrade, or use an offer = PROMOTIONS."),
    ("UPDATES/PROMOTIONS examples", "Synthetic: a cloud subscription renewal confirmation = UPDATES. An email urging an upgrade with a 30% discount = PROMOTIONS."),
    ("PERSONAL vs OTHER", "A direct message involving a recognizable personal relationship/arrangement = PERSONAL. Otherwise use the best-fitting category or OTHER."),
    ("PERSONAL/OTHER examples", "Synthetic: a friend asking about Saturday dinner = PERSONAL. A public residents association agenda notice = OTHER."),
    ("Response/priority examples", "Synthetic: confirm an interview slot by 5 PM today = needs_response=yes, URGENT. An automated active-security notice requiring immediate password reset but no reply = needs_response=no, URGENT."),
    ("When unsure", "Set unsure=yes, explain briefly in note, and leave only genuinely uncertain task fields blank. Never guess or consult model outputs/old labels."),
    ("Body preview", "The workbook contains a configurable body preview (default 1,500 characters). The full body remains in the local database and can be recovered using email_id; do not copy full bodies into documentation or reports."),
    ("Pilot and freeze", "Label 20 emails first. Review difficult cases once, refine the guide once, then freeze a final version. Any later semantic change requires a new guide version, changelog entry, and identification/relabeling of affected rows."),
]


class WorkbookExportError(ValueError):
    """Safe-to-display workbook export error; never includes email content or IDs."""


def _clean_text(value: Any, *, limit: int = MAX_CELL_CHARS) -> str:
    text = "" if value is None else str(value)
    text = XML_ILLEGAL.sub("", text)
    return text[: min(limit, MAX_CELL_CHARS)]


def _write_text(cell: Any, value: Any, *, limit: int = MAX_CELL_CHARS, text_format: bool = False) -> None:
    """Write untrusted content as an explicit string, never an Excel formula."""
    cell.value = _clean_text(value, limit=limit)
    cell.data_type = "s"
    if text_format:
        cell.number_format = "@"


def read_source_csv(path: Path) -> list[dict[str, str]]:
    if not path.is_file():
        raise WorkbookExportError("Source dataset not found. Export the local email dataset first.")
    try:
        with path.open("r", newline="", encoding="utf-8-sig") as handle:
            reader = csv.DictReader(handle)
            if not reader.fieldnames or "email_id" not in reader.fieldnames:
                raise WorkbookExportError("Source dataset must contain an email_id column.")
            rows = [{key: (value or "") for key, value in row.items() if key is not None} for row in reader]
    except (OSError, csv.Error) as exc:
        raise WorkbookExportError(f"Could not read source dataset ({type(exc).__name__}).") from exc
    ids = [(row.get("email_id") or "").strip() for row in rows]
    if any(not item for item in ids) or len(ids) != len(set(ids)):
        raise WorkbookExportError("Source dataset contains missing or duplicate email identifiers.")
    required_context = {"sender", "subject", "snippet", "body_text"}
    if not required_context.issubset(set(reader.fieldnames or [])):
        raise WorkbookExportError("Source dataset is missing required email context columns.")
    return rows


def _is_previously_verified(row: dict[str, str]) -> bool:
    """Use the existing labeler's completeness convention for blind relabeling."""
    return _fully_labeled(row)


def _read_queue_ids(path: Path) -> set[str]:
    try:
        return _load_queue_ids(path)
    except (FileNotFoundError, ValueError) as exc:
        raise WorkbookExportError("Minority labeling queue is missing or invalid.") from exc


def _read_frozen_ids(path: Path) -> set[str]:
    try:
        return read_frozen_test_ids(path)
    except (FileNotFoundError, ValueError) as exc:
        raise WorkbookExportError("Existing frozen split manifest is missing or invalid.") from exc


def _test_v2_ids(
    rows: list[dict[str, str]], *, split_manifest: Path, test_manifest: Path,
    seed: int, sample_size: int = 60,
) -> tuple[list[str], bool]:
    frozen = _read_frozen_ids(split_manifest)
    by_id = {(row.get("email_id") or "").strip(): row for row in rows}
    if test_manifest.exists():
        try:
            manifest = json.loads(test_manifest.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError) as exc:
            raise WorkbookExportError("Existing test_v2 manifest is unreadable.") from exc
        selected = manifest.get("email_ids") if isinstance(manifest, dict) else None
        if manifest.get("version") != 1 or not isinstance(selected, list):
            raise WorkbookExportError("Existing test_v2 manifest has an unsupported schema.")
        if len(selected) != sample_size or len(set(selected)) != len(selected) or not all(isinstance(x, str) for x in selected):
            raise WorkbookExportError("Existing test_v2 manifest does not match the required 60 unique IDs.")
        if set(selected) & frozen or set(selected) - set(by_id):
            raise WorkbookExportError("Existing test_v2 manifest conflicts with the source or frozen split.")
        # On resume, prior human labels are expected: the frozen test membership must persist.
        return list(selected), False

    eligible = [
        email_id for email_id, row in by_id.items()
        if email_id not in frozen and not has_any_human_label(row)
    ]
    if len(eligible) < sample_size:
        raise WorkbookExportError("Fewer than 60 never-before-labeled, non-frozen emails are available.")
    selected = random.Random(seed).sample(sorted(eligible), sample_size)
    test_manifest.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "version": 1,
        "sample_name": "test_v2",
        "seed": seed,
        "sample_size": sample_size,
        "email_ids": selected,
        "policy": "Frozen blind test sample. Never use these labels for training, tuning, or model selection.",
    }
    _atomic_write_json_new(test_manifest, payload)
    return selected, True


def _atomic_write_json_new(path: Path, payload: dict[str, Any]) -> None:
    if path.exists():
        raise WorkbookExportError("Refusing to overwrite an existing test_v2 manifest.")
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_suffix(path.suffix + ".tmp")
    try:
        with temp.open("x", encoding="utf-8") as handle:
            json.dump(payload, handle, indent=2)
            handle.write("\n")
        try:
            temp.replace(path)
        except OSError:
            if path.exists():
                raise WorkbookExportError("Test manifest appeared during export; refusing to overwrite it.")
            raise
    finally:
        if temp.exists():
            temp.unlink()


def _select_ids(
    rows: list[dict[str, str]], *, mode: str, seed: int, count: int | None,
    queue_path: Path, split_manifest: Path, test_manifest: Path,
) -> tuple[list[str], str]:
    by_id = {(row.get("email_id") or "").strip(): row for row in rows}
    if mode == "relabel":
        candidates = [email_id for email_id, row in by_id.items() if _is_previously_verified(row)]
        purpose = "BLIND RELABEL — old labels are hidden; do not consult previous annotations."
    elif mode == "minority":
        queued = _read_queue_ids(queue_path)
        candidates = [email_id for email_id in queued if email_id in by_id and not has_any_human_label(by_id[email_id])]
        purpose = "MINORITY CLASS CANDIDATES — heuristic queue membership is not a label."
    elif mode == "test_v2":
        candidates, _ = _test_v2_ids(rows, split_manifest=split_manifest, test_manifest=test_manifest, seed=seed)
        purpose = "TEST — frozen blind holdout. Never use these labels for training, tuning, or model selection."
    else:
        candidates = [email_id for email_id, row in by_id.items() if not has_any_human_label(row)]
        if count is None or count < 1:
            raise WorkbookExportError("random mode requires --count N with N greater than zero.")
        if len(candidates) < count:
            raise WorkbookExportError("Not enough never-before-labeled emails for the requested random count.")
        candidates = random.Random(seed).sample(sorted(candidates), count)
        purpose = "RANDOM UNLABELED SAMPLE."
    if not candidates:
        raise WorkbookExportError("No eligible rows were found for the requested mode.")
    # A seeded shuffle prevents chronological or queue-order labeling blocks.
    random.Random(seed).shuffle(candidates)
    return candidates, purpose


def _add_defined_name(workbook: Workbook, name: str, sheet: str, cell_range: str) -> None:
    ref = f"'{sheet}'!{cell_range}"
    defined = DefinedName(name, attr_text=ref)
    try:
        workbook.defined_names.add(defined)
    except AttributeError:  # openpyxl 3.0 compatibility
        workbook.defined_names.append(defined)


def _backup_existing(path: Path, backup_dir: Path) -> Path:
    backup_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S%fZ")
    backup = backup_dir / f"{path.stem}.{stamp}{path.suffix}"
    shutil.copy2(path, backup)
    return backup


def _write_workbook(
    output_path: Path, selected_rows: list[dict[str, str]], *, purpose: str,
    mode: str, batch_name: str, body_preview_chars: int, schema_version: str = SCHEMA_VERSION,
) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Labeling"
    guide = wb.create_sheet("Guide")
    progress = wb.create_sheet("Progress")
    lists = wb.create_sheet("Lists")

    headers = list(VISIBLE_COLUMNS)
    ws.append(headers)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.protection = Protection(locked=True)

    for row_number, source in enumerate(selected_rows, start=1):
        values = [
            row_number,
            source.get("email_id", ""),
            source.get("sender", ""),
            source.get("subject", ""),
            source.get("snippet", ""),
            _clean_text(source.get("body_text") or source.get("snippet") or "")[:body_preview_chars],
            "", "", "", "", "no", "",
        ]
        # Append an empty row first: raw untrusted text may contain XML-illegal
        # characters or formula prefixes, so sanitize/force-string before assignment.
        ws.append([row_number] + [None] * (len(headers) - 1))
        excel_row = ws.max_row
        for col in range(2, len(headers) + 1):
            cell = ws.cell(excel_row, col)
            _write_text(cell, values[col - 1], limit=body_preview_chars if col == 6 else MAX_CELL_CHARS, text_format=(col == 2))
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.protection = Protection(locked=col not in range(7, 13))
        ws.cell(excel_row, 1).alignment = Alignment(vertical="top", horizontal="center")
        # Dynamic but bounded row heights keep the body preview readable.
        preview = str(values[5])
        estimated_lines = max(3, min(12, (len(preview) // 95) + preview.count("\n") + 1))
        ws.row_dimensions[excel_row].height = min(180, max(60, estimated_lines * 15))

    widths = {"A": 8, "B": 23, "C": 30, "D": 34, "E": 38, "F": 65,
              "G": 14, "H": 19, "I": 20, "J": 18, "K": 12, "L": 38}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:L{ws.max_row}"
    ws.sheet_view.showGridLines = False
    ws.protection.sheet = True
    # In OOXML these are protection flags: True DISABLES selection.
    # Keep both False so Excel users can select unlocked label cells
    # (and can still navigate locked context cells if desired).
    ws.protection.selectLockedCells = False
    ws.protection.selectUnlockedCells = False

    # Hidden dropdown source values and workbook-level names.
    lists.append(["yes_no", "priority", "category"])
    for idx, value in enumerate(YES_NO, start=2):
        lists.cell(idx, 1, value)
    for idx, value in enumerate(PRIORITIES, start=2):
        lists.cell(idx, 2, value)
    for idx, value in enumerate(CATEGORIES, start=2):
        lists.cell(idx, 3, value)
    lists.sheet_state = "hidden"
    _add_defined_name(wb, "YesNo", "Lists", "$A$2:$A$3")
    _add_defined_name(wb, "PriorityValues", "Lists", "$B$2:$B$4")
    _add_defined_name(wb, "CategoryValues", "Lists", "$C$2:$C$7")

    dropdowns = (("G", "=YesNo"), ("H", "=YesNo"), ("I", "=PriorityValues"),
                 ("J", "=CategoryValues"), ("K", "=YesNo"))
    for column, formula in dropdowns:
        validation = DataValidation(type="list", formula1=formula, allow_blank=True)
        validation.errorTitle = "Invalid label"
        validation.error = "Choose a value from the dropdown list."
        validation.promptTitle = "Labeling guide"
        validation.prompt = "Use the Guide sheet before selecting a label."
        validation.showErrorMessage = True
        validation.showInputMessage = True
        ws.add_data_validation(validation)
        validation.add(f"{column}2:{column}{max(ws.max_row, 2)}")

    blank_fill = PatternFill("solid", fgColor="FFF2CC")
    for column in "GHIJ":
        ws.conditional_formatting.add(
            f"{column}2:{column}{max(ws.max_row, 2)}",
            FormulaRule(formula=[f'LEN(TRIM({column}2&""))=0'], fill=blank_fill),
        )
    conflict_fill = PatternFill("solid", fgColor="F4CCCC")
    if ws.max_row >= 2:
        ws.conditional_formatting.add(
            f"H2:I{ws.max_row}",
            FormulaRule(
                formula=['OR(AND($H2="yes",$I2="INFORMATIONAL"),AND($H2="no",$I2="NEEDS_REPLY"))'],
                fill=conflict_fill,
            ),
        )

    guide.append(["Topic", "Definition and decision rule"])
    for cell in guide[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
    for item in GUIDE_ROWS:
        guide.append(list(item))
    guide.column_dimensions["A"].width = 25
    guide.column_dimensions["B"].width = 115
    guide.freeze_panes = "A2"
    for row in guide.iter_rows(min_row=2):
        for cell in row:
            _write_text(cell, cell.value)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        guide.row_dimensions[row[0].row].height = 42

    progress.append(["Metric", "Value"])
    progress.append(["Batch name", batch_name])
    progress.append(["Mode", mode])
    progress.append(["Purpose", purpose])
    progress.append(["Schema version", schema_version])
    progress.append(["Guide version", GUIDE_VERSION])
    progress.append(["Rows in this workbook", f"=COUNTA(Labeling!$A$2:$A${ws.max_row})"])
    progress.append(["Fully labeled rows", f'=COUNTIFS(Labeling!$G$2:$G${ws.max_row},"<>",Labeling!$H$2:$H${ws.max_row},"<>",Labeling!$I$2:$I${ws.max_row},"<>",Labeling!$J$2:$J${ws.max_row},"<>")'])
    progress.append(["Remaining rows", "=B7-B8"])
    progress.append(["Uncertain rows", f'=COUNTIF(Labeling!$K$2:$K${ws.max_row},"yes")'])
    progress.append(["", ""])
    progress.append(["Label field", "Value", "Count"])
    for cell in progress[12]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
    for field, column, values in (
        ("is_spam", "G", YES_NO), ("needs_response", "H", YES_NO),
        ("priority", "I", PRIORITIES), ("category", "J", CATEGORIES),
    ):
        for value in values:
            progress.append([field, value, f'=COUNTIF(Labeling!${column}$2:${column}${ws.max_row},B{progress.max_row + 1})'])
    progress.column_dimensions["A"].width = 28
    progress.column_dimensions["B"].width = 75
    progress.column_dimensions["C"].width = 16
    progress.freeze_panes = "A2"
    for row in progress.iter_rows():
        for cell in row:
            if cell.data_type == "s":
                _write_text(cell, cell.value)
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    wb.properties.title = "FRIDAY Blind Email Labeling Workbook"
    wb.properties.subject = f"Blind {mode} labeling batch; schema {schema_version}"
    wb.properties.creator = "FRIDAY local labeling workflow"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path = output_path.with_name(f"{output_path.stem}.tmp{output_path.suffix}")
    try:
        wb.save(temporary_path)
        temporary_path.replace(output_path)
    finally:
        if temporary_path.exists():
            temporary_path.unlink()


def export_labeling_workbook(
    input_path: Path = DEFAULT_INPUT, *, mode: str, batch_name: str,
    output_dir: Path = DEFAULT_OUTPUT_DIR, seed: int = DEFAULT_SEED,
    batch_size: int = DEFAULT_BATCH_SIZE, offset: int = 0, count: int | None = None,
    body_preview_chars: int = DEFAULT_BODY_PREVIEW_CHARS,
    queue_path: Path = DEFAULT_QUEUE, split_manifest: Path = DEFAULT_SPLIT_MANIFEST,
    test_manifest: Path = DEFAULT_TEST_MANIFEST, overwrite: bool = False,
) -> dict[str, Any]:
    if mode not in {"relabel", "minority", "test_v2", "random"}:
        raise WorkbookExportError("Unsupported selection mode.")
    if batch_size < 1 or offset < 0 or body_preview_chars < 1:
        raise WorkbookExportError("batch size and preview length must be positive; offset cannot be negative.")
    if not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_.-]{0,79}", batch_name):
        raise WorkbookExportError("Batch name must contain only letters, numbers, dots, underscores, or hyphens.")

    rows = read_source_csv(input_path)
    ids, purpose = _select_ids(
        rows, mode=mode, seed=seed, count=count, queue_path=queue_path,
        split_manifest=split_manifest, test_manifest=test_manifest,
    )
    total_candidates = len(ids)
    selected_ids = ids[offset: offset + batch_size]
    if not selected_ids:
        raise WorkbookExportError("No rows remain at this batch offset.")
    by_id = {(row.get("email_id") or "").strip(): row for row in rows}
    selected_rows = [by_id[email_id] for email_id in selected_ids]
    output_path = output_dir / f"{batch_name}.xlsx"
    if output_path.exists():
        if not overwrite:
            raise WorkbookExportError("Output workbook already exists. Use --overwrite to replace it; a backup will be created.")
        _backup_existing(output_path, output_dir.parent / "backups")
    _write_workbook(
        output_path, selected_rows, purpose=purpose, mode=mode,
        batch_name=batch_name, body_preview_chars=body_preview_chars,
    )
    return {
        "mode": mode,
        "batch_name": batch_name,
        "rows_in_workbook": len(selected_rows),
        "eligible_rows": total_candidates,
        "remaining_after_batch": max(0, total_candidates - offset - len(selected_rows)),
        "seed": seed,
        "offset": offset,
        "output_file": output_path.name,
        "test_batch": mode == "test_v2",
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--mode", choices=("relabel", "minority", "test_v2", "random"), required=True)
    parser.add_argument("--batch-name", required=True)
    parser.add_argument("--seed", type=int, default=DEFAULT_SEED)
    parser.add_argument("--batch-size", type=int, default=DEFAULT_BATCH_SIZE)
    parser.add_argument("--offset", type=int, default=0, help="Start position in the deterministic shuffled selection.")
    parser.add_argument("--count", type=int, help="Required for --mode random; number of never-labeled emails to sample.")
    parser.add_argument("--body-preview-chars", type=int, default=DEFAULT_BODY_PREVIEW_CHARS)
    parser.add_argument("--queue", type=Path, default=DEFAULT_QUEUE)
    parser.add_argument("--split-manifest", type=Path, default=DEFAULT_SPLIT_MANIFEST)
    parser.add_argument("--test-manifest", type=Path, default=DEFAULT_TEST_MANIFEST)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--overwrite", action="store_true")
    args = parser.parse_args()
    try:
        result = export_labeling_workbook(
            args.input, mode=args.mode, batch_name=args.batch_name,
            output_dir=args.output_dir, seed=args.seed, batch_size=args.batch_size,
            offset=args.offset, count=args.count, body_preview_chars=args.body_preview_chars,
            queue_path=args.queue, split_manifest=args.split_manifest,
            test_manifest=args.test_manifest, overwrite=args.overwrite,
        )
    except (WorkbookExportError, OSError) as exc:
        parser.error(str(exc))
    print(json.dumps(result, indent=2))
    print("Private workbook created locally. Console output contains aggregate counts only.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
