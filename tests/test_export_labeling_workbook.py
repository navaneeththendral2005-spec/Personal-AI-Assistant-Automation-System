from __future__ import annotations

import csv
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from scripts.export_labeling_workbook import (
    LABEL_COLUMNS,
    VISIBLE_COLUMNS,
    WorkbookExportError,
    export_labeling_workbook,
)


def make_rows(count: int = 8) -> list[dict[str, str]]:
    rows = []
    for index in range(count):
        rows.append({
            "email_id": f"message-{index:03d}",
            "sender": f"sender{index}@example.test",
            "subject": f"Synthetic subject {index}",
            "snippet": f"Synthetic snippet {index}",
            "body_text": f"Synthetic body {index}",
            "labels": "INBOX",
            "true_is_spam": "",
            "true_priority": "",
            "true_needs_response": "",
            "true_category": "",
        })
    return rows


def write_source(path: Path, rows: list[dict[str, str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fields = list(rows[0])
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def write_split(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps({"version": 1, "tasks": {"priority": {"test_ids": ["frozen-1"]}}}), encoding="utf-8")


def test_blind_workbook_has_only_allowed_columns_and_hides_old_labels(tmp_path: Path) -> None:
    rows = make_rows(3)
    rows[0].update({
        "true_is_spam": "no", "true_priority": "URGENT",
        "true_needs_response": "yes", "true_category": "JOB",
        "predicted_priority": "INFORMATIONAL", "rules_category": "OTHER",
        "llm_is_spam": "yes", "label_source": "model",
    })
    source = tmp_path / "email_labels.csv"
    write_source(source, rows)
    output_dir = tmp_path / "labeling"
    result = export_labeling_workbook(
        source, mode="relabel", batch_name="blind", output_dir=output_dir,
        batch_size=10, seed=17,
    )
    assert result["rows_in_workbook"] == 1

    workbook = load_workbook(output_dir / "blind.xlsx", data_only=False)
    sheet = workbook["Labeling"]
    headers = [cell.value for cell in sheet[1]]
    assert tuple(headers) == VISIBLE_COLUMNS
    assert not any(
        str(header).startswith(("true_", "final_", "predicted_", "rules_", "llm_", "review"))
        or str(header) == "label_source"
        for header in headers
    )
    rendered = " ".join(str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value is not None)
    assert "URGENT" not in rendered
    assert "INFORMATIONAL" not in rendered
    assert "model" not in rendered
    assert workbook["Lists"].sheet_state == "hidden"
    assert sheet.protection.sheet is True
    # OOXML sheet-protection flags are inverted: True blocks selection.
    assert sheet.protection.selectLockedCells is False
    assert sheet.protection.selectUnlockedCells is False
    for col in range(1, 7):
        assert sheet.cell(2, col).protection.locked is True
    for col in range(7, 13):
        assert sheet.cell(2, col).protection.locked is False
    assert len(sheet.data_validations.dataValidation) == 5
    assert sheet.freeze_panes == "C2"
    assert sheet.auto_filter.ref.startswith("A1:L")
    assert len(sheet.conditional_formatting) >= 2
    assert "COUNTIF" in workbook["Progress"]["C13"].value


def test_formula_like_body_and_scientific_notation_id_round_trip_as_text(tmp_path: Path) -> None:
    rows = make_rows(1)
    rows[0]["email_id"] = "1e5ABC-987654321"
    rows[0]["body_text"] = '=HYPERLINK("https://example.test","click")'
    source = tmp_path / "source.csv"
    write_source(source, rows)
    out = tmp_path / "labeling"
    export_labeling_workbook(source, mode="random", batch_name="safe", output_dir=out, count=1)

    workbook = load_workbook(out / "safe.xlsx", data_only=False)
    sheet = workbook["Labeling"]
    assert sheet["B2"].value == "1e5ABC-987654321"
    assert sheet["B2"].data_type == "s"
    assert sheet["B2"].number_format == "@"
    assert sheet["F2"].value == '=HYPERLINK("https://example.test","click")'
    assert sheet["F2"].data_type == "s"


def test_body_preview_truncates_and_illegal_xml_characters_are_removed(tmp_path: Path) -> None:
    rows = make_rows(1)
    rows[0]["body_text"] = "A\x00B" + "x" * 2000
    source = tmp_path / "source.csv"
    write_source(source, rows)
    out = tmp_path / "labeling"
    export_labeling_workbook(source, mode="random", batch_name="preview", output_dir=out, count=1, body_preview_chars=20)
    sheet = load_workbook(out / "preview.xlsx")["Labeling"]
    assert sheet["F2"].value == "AB" + "x" * 18
    assert len(sheet["F2"].value) == 20


def test_seeded_random_order_is_reproducible(tmp_path: Path) -> None:
    source = tmp_path / "source.csv"
    write_source(source, make_rows(8))
    out = tmp_path / "labeling"
    export_labeling_workbook(source, mode="random", batch_name="one", output_dir=out, count=6, seed=123)
    export_labeling_workbook(source, mode="random", batch_name="two", output_dir=out, count=6, seed=123)
    one = [r[0] for r in load_workbook(out / "one.xlsx")["Labeling"].iter_rows(min_row=2, min_col=2, max_col=2, values_only=True)]
    two = [r[0] for r in load_workbook(out / "two.xlsx")["Labeling"].iter_rows(min_row=2, min_col=2, max_col=2, values_only=True)]
    assert one == two


def test_random_mode_requires_count_and_only_selects_unlabeled(tmp_path: Path) -> None:
    rows = make_rows(4)
    rows[0]["true_category"] = "JOB"
    source = tmp_path / "source.csv"
    write_source(source, rows)
    with pytest.raises(WorkbookExportError, match="requires --count"):
        export_labeling_workbook(source, mode="random", batch_name="missing-count", output_dir=tmp_path / "out")
    result = export_labeling_workbook(source, mode="random", batch_name="random", output_dir=tmp_path / "out", count=3)
    assert result["eligible_rows"] == 3


def test_minority_mode_uses_queue_ids_without_exposing_queue_metadata(tmp_path: Path) -> None:
    source = tmp_path / "source.csv"
    write_source(source, make_rows(4))
    queue = tmp_path / "queue.csv"
    queue.write_text("email_id,focus_class,score\nmessage-001,FINANCE,9.2\nmessage-003,OTHER,7.1\n", encoding="utf-8")
    out = tmp_path / "labeling"
    export_labeling_workbook(source, mode="minority", batch_name="minority", output_dir=out, queue_path=queue)
    sheet = load_workbook(out / "minority.xlsx")["Labeling"]
    assert sheet.max_row == 3
    assert "focus_class" not in [cell.value for cell in sheet[1]]
    assert "score" not in [cell.value for cell in sheet[1]]


def test_test_v2_selects_60_unlabeled_ids_and_freezes_manifest(tmp_path: Path) -> None:
    rows = make_rows(65)
    rows[0]["true_category"] = "JOB"
    rows[1]["true_is_spam"] = "no"
    source = tmp_path / "source.csv"
    write_source(source, rows)
    split = tmp_path / "split.json"
    write_split(split)
    test_manifest = tmp_path / "test_v2_manifest.json"
    out = tmp_path / "labeling"
    result = export_labeling_workbook(
        source, mode="test_v2", batch_name="test-a", output_dir=out,
        split_manifest=split, test_manifest=test_manifest, batch_size=40, seed=44,
    )
    assert result["eligible_rows"] == 60
    assert result["rows_in_workbook"] == 40
    manifest = json.loads(test_manifest.read_text(encoding="utf-8"))
    assert len(manifest["email_ids"]) == 60
    assert not ({"message-000", "message-001", "frozen-1"} & set(manifest["email_ids"]))
    assert result["test_batch"] is True
    # Resume with offset 40 uses the same frozen sample and supplies the remaining 20.
    result2 = export_labeling_workbook(
        source, mode="test_v2", batch_name="test-b", output_dir=out,
        split_manifest=split, test_manifest=test_manifest, batch_size=40, offset=40, seed=999,
    )
    assert result2["rows_in_workbook"] == 20


def test_output_refuses_overwrite_then_backs_up_when_explicitly_allowed(tmp_path: Path) -> None:
    source = tmp_path / "source.csv"
    write_source(source, make_rows(2))
    out = tmp_path / "labeling"
    export_labeling_workbook(source, mode="random", batch_name="batch", output_dir=out, count=1)
    with pytest.raises(WorkbookExportError, match="already exists"):
        export_labeling_workbook(source, mode="random", batch_name="batch", output_dir=out, count=1)
    export_labeling_workbook(source, mode="random", batch_name="batch", output_dir=out, count=1, overwrite=True)
    backups = list((tmp_path / "backups").glob("batch.*.xlsx"))
    assert len(backups) == 1
    assert load_workbook(out / "batch.xlsx").sheetnames == ["Labeling", "Guide", "Progress", "Lists"]


def test_batch_offset_and_batch_size_are_applied(tmp_path: Path) -> None:
    source = tmp_path / "source.csv"
    write_source(source, make_rows(8))
    out = tmp_path / "labeling"
    first = export_labeling_workbook(source, mode="random", batch_name="part1", output_dir=out, count=7, batch_size=3, offset=0, seed=9)
    second = export_labeling_workbook(source, mode="random", batch_name="part2", output_dir=out, count=7, batch_size=3, offset=3, seed=9)
    assert first["rows_in_workbook"] == 3
    assert second["rows_in_workbook"] == 3
    ids1 = {r[0] for r in load_workbook(out / "part1.xlsx")["Labeling"].iter_rows(min_row=2, min_col=2, max_col=2, values_only=True)}
    ids2 = {r[0] for r in load_workbook(out / "part2.xlsx")["Labeling"].iter_rows(min_row=2, min_col=2, max_col=2, values_only=True)}
    assert ids1.isdisjoint(ids2)
